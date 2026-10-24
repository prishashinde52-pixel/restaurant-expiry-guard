from __future__ import annotations

from datetime import date, datetime
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path
import os
import re
import smtplib
import sys
import tempfile
import uuid

import pandas as pd
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image, ImageEnhance, ImageOps
import streamlit as st

ROOT = Path(__file__).resolve().parent
DATA_FILE = ROOT / "data" / "inventory.xlsx"
HEADERS = [
    "Batch ID", "Ingredient", "Category", "Received", "Expiry date",
    "Date type", "Qty received", "Qty remaining", "Unit", "Supplier",
    "Storage", "State", "Notes",
]
DATE_PATTERN = re.compile(
    r"\b(?:\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})\b"
)
EXPIRY_WORDS = ("exp", "use by", "use-by", "best before", "bb", "expiry", "expires")
MANUFACTURE_WORDS = ("mfg", "mfd", "manufactured", "packed on", "pkd")
COLOURS = {"EXPIRED": "#ffe0e0", "PAST BEST BEFORE": "#fff0cb", "EXPIRING SOON": "#fff0cb", "OK": "#e4f5e8", "CLOSED": "#eeeeee"}


def clean_text(text: object) -> str:
    """Keep untrusted user / OCR text from becoming Excel formulas."""
    value = str(text or "").strip()[:500]
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def save_book(workbook: Workbook) -> None:
    """Write a temporary workbook first, then replace the original."""
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_path = tempfile.mkstemp(suffix=".xlsx", dir=DATA_FILE.parent)
    os.close(fd)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, DATA_FILE)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def ensure_book() -> None:
    if DATA_FILE.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(HEADERS)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:M1"
    ws.sheet_view.showGridLines = False
    widths = [15, 26, 18, 17, 17, 20, 16, 16, 12, 22, 18, 13, 36]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")

    settings = wb.create_sheet("Settings")
    settings["A1"] = "Days before expiry to warn"
    settings["B1"] = 3
    settings.column_dimensions["A"].width = 34
    wb.defined_names.add(DefinedName("AlertDays", attr_text="'Settings'!$B$1"))

    movements = wb.create_sheet("Movements")
    movements.append(["Recorded at", "Batch ID", "Ingredient", "Action", "Quantity", "Remaining", "Notes"])
    for cell in movements[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E777D")
    for col in "ABCDEFG":
        movements.column_dimensions[col].width = 22
    movements.column_dimensions["G"].width = 40

    # Excel recalculates these colours as TODAY() changes.
    rules = [
        ('AND($L2="Active",$H2>0,$F2<>"Best before",$E2<TODAY())', "FFE0E0"),
        ('AND($L2="Active",$H2>0,OR(AND($F2="Best before",$E2<TODAY()),AND($E2>=TODAY(),$E2<=TODAY()+AlertDays)))', "FFF0CB"),
        ('AND($L2="Active",$H2>0,$E2>TODAY()+AlertDays)', "E4F5E8"),
        ('OR($L2="Used",$L2="Discarded")', "EEEEEE"),
    ]
    for formula, color in rules:
        ws.conditional_formatting.add(
            "A2:M10000", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color))
        )
    save_book(wb)
    wb.close()


def alert_days() -> int:
    ensure_book()
    wb = load_workbook(DATA_FILE, read_only=True)
    try:
        return max(0, int(wb["Settings"]["B1"].value or 0))
    finally:
        wb.close()


def set_alert_days(days: int) -> None:
    if not 0 <= days <= 90:
        raise ValueError("Alert window must be 0-90 days.")
    wb = load_workbook(DATA_FILE)
    wb["Settings"]["B1"] = days
    save_book(wb)
    wb.close()


def read_batches() -> list[dict]:
    ensure_book()
    wb = load_workbook(DATA_FILE, read_only=True, data_only=True)
    try:
        items = []
        for cells in wb["Inventory"].iter_rows(min_row=2, values_only=True):
            if not cells[0]:
                continue
            item = dict(zip(HEADERS, cells))
            for column in ("Received", "Expiry date"):
                value = item[column]
                if isinstance(value, datetime):
                    item[column] = value.date()
            items.append(item)
        return items
    finally:
        wb.close()


def batch_status(item: dict, days: int, today: date | None = None) -> str:
    today = today or date.today()
    if item["State"] != "Active" or float(item["Qty remaining"] or 0) <= 0:
        return "CLOSED"
    expiry = item["Expiry date"]
    if expiry < today:
        return "PAST BEST BEFORE" if item["Date type"] == "Best before" else "EXPIRED"
    if (expiry - today).days <= days:
        return "EXPIRING SOON"
    return "OK"


def add_batch(name: str, category: str, received: date, expiry: date,
              date_type: str, quantity: float, unit: str, supplier: str,
              storage: str, notes: str) -> str:
    if not name.strip() or quantity <= 0:
        raise ValueError("Ingredient name and a positive quantity are required.")
    ensure_book()
    batch_id = uuid.uuid4().hex[:10].upper()
    wb = load_workbook(DATA_FILE)
    ws = wb["Inventory"]
    ws.append([
        batch_id, clean_text(name), clean_text(category), received, expiry,
        date_type, float(quantity), float(quantity), clean_text(unit),
        clean_text(supplier), clean_text(storage), "Active", clean_text(notes),
    ])
    row = ws.max_row
    ws.cell(row, 4).number_format = "DD-MMM-YYYY"
    ws.cell(row, 5).number_format = "DD-MMM-YYYY"
    save_book(wb)
    wb.close()
    return batch_id


def record_movement(batch_id: str, action: str, quantity: float, note: str = "") -> float:
    if action not in ("Used", "Discarded") or quantity <= 0:
        raise ValueError("Choose Used or Discarded and a positive quantity.")
    wb = load_workbook(DATA_FILE)
    try:
        ws = wb["Inventory"]
        for row in range(2, ws.max_row + 1):
            if ws.cell(row, 1).value != batch_id:
                continue
            if ws.cell(row, 12).value != "Active":
                raise ValueError("This batch is already closed.")
            old = float(ws.cell(row, 8).value or 0)
            if quantity > old + 0.000001:
                raise ValueError(f"Only {old:g} units remain in this batch.")
            remaining = max(0.0, round(old - quantity, 3))
            ws.cell(row, 8).value = remaining
            if remaining == 0:
                ws.cell(row, 12).value = action
            wb["Movements"].append([
                datetime.now(), batch_id, ws.cell(row, 2).value,
                action, float(quantity), remaining, clean_text(note),
            ])
            save_book(wb)
            return remaining
        raise ValueError("Batch ID not found.")
    finally:
        wb.close()


def parse_date(value: str) -> date | None:
    """Accept YYYY-MM-DD, Indian DD/MM/YYYY, or unambiguous US MM/DD/YYYY."""
    parts = re.split(r"[-/.]", value)
    if len(parts) != 3:
        return None
    try:
        a, b, c = map(int, parts)
        if len(parts[0]) == 4:
            year, month, day = a, b, c
        else:
            year = c + 2000 if c < 100 else c
            # For ambiguous labels, assume DD/MM/YYYY; ask for confirmation.
            day, month = a, b
            if a <= 12 and b > 12:  # Unambiguously MM/DD/YYYY
                month, day = a, b
        result = date(year, month, day)
        return result if 2020 <= year <= 2100 else None
    except ValueError:
        return None


def photo_date_suggestions(image_bytes: bytes) -> tuple[str, list[date]]:
    """Use local Tesseract OCR; return only suggestions, never auto-save."""
    import pytesseract  # OCR remains optional for manual-entry users.

    picture = Image.open(BytesIO(image_bytes)).convert("RGB")
    picture = ImageEnhance.Contrast(ImageOps.grayscale(picture)).enhance(2)
    ocr_text = pytesseract.image_to_string(picture)
    ranked: list[tuple[int, date]] = []
    for line in ocr_text.splitlines():
        lower = line.lower()
        has_expiry_word = any(word in lower for word in EXPIRY_WORDS)
        if any(word in lower for word in MANUFACTURE_WORDS) and not has_expiry_word:
            continue
        for match in DATE_PATTERN.finditer(line):
            found = parse_date(match.group())
            if found is not None:
                ranked.append((0 if has_expiry_word else 1, found))
    suggestions = []
    for _, parsed in sorted(ranked, key=lambda pair: pair[0]):
        if parsed not in suggestions:
            suggestions.append(parsed)
    return ocr_text, suggestions


def send_email_alerts() -> str:
    """Email due batches. Run manually or schedule python app.py --check-alerts."""
    load_dotenv(ROOT / ".env")
    host = os.getenv("SMTP_HOST", "smtp.gmail.com")
    port = int(os.getenv("SMTP_PORT", "587"))
    username = os.getenv("SMTP_USER", "")
    password = os.getenv("SMTP_PASSWORD", "")
    recipient = os.getenv("ALERT_TO", "")
    if not username or not password or not recipient:
        raise ValueError("Create a .env file with SMTP_USER, SMTP_PASSWORD and ALERT_TO.")

    days = alert_days()
    due = [item for item in read_batches() if batch_status(item, days) in ("EXPIRED", "PAST BEST BEFORE", "EXPIRING SOON")]
    if not due:
        return "No expiring or expired batches; no email needed."

    lines = [f"Restaurant inventory alert ({date.today().isoformat()})", ""]
    for item in sorted(due, key=lambda i: i["Expiry date"]):
        lines.append(
            f"{batch_status(item, days)} | {item['Ingredient']} | "
            f"{item['Qty remaining']:g} {item['Unit']} | "
            f"Date: {item['Expiry date'].isoformat()} | Batch: {item['Batch ID']}"
        )
    lines += ["", "Check labels and storage conditions. Do not use expired items without following food-safety rules."]
    mail = EmailMessage()
    mail["Subject"] = f"Restaurant Expiry Guard: {len(due)} batch(es) need attention"
    mail["From"] = username
    mail["To"] = recipient
    mail.set_content("\n".join(lines))

    if port == 465:
        import ssl
        with smtplib.SMTP_SSL(host, port, context=ssl.create_default_context(), timeout=20) as server:
            server.login(username, password)
            server.send_message(mail)
    else:
        with smtplib.SMTP(host, port, timeout=20) as server:
            server.starttls()
            server.login(username, password)
            server.send_message(mail)
    return f"Email sent to {recipient} for {len(due)} batch(es)."


def main() -> None:
    st.set_page_config(page_title="Restaurant Expiry Guard",  layout="wide")
    ensure_book()
    days = alert_days()
    items = read_batches()

    st.title(" Restaurant Expiry Guard")
    st.caption("Track delivered ingredients, confirm expiry labels and reduce food waste.")
    dashboard, intake, usage, settings = st.tabs([
        "Dashboard", " Add stock", " Used / discarded", " Settings & alerts"
    ])

    with dashboard:
        statuses = [batch_status(item, days) for item in items]
        columns = st.columns(4)
        for column, title, amount in zip(
            columns,
            ["In-stock batches", "Expiring soon", "Past labelled date", "Closed batches"],
            [len(statuses) - statuses.count("CLOSED"), statuses.count("EXPIRING SOON"),
             statuses.count("EXPIRED") + statuses.count("PAST BEST BEFORE"), statuses.count("CLOSED")],
        ):
            column.metric(title, amount)

        include_closed = st.checkbox("Include used / discarded batches", value=False)
        visible = []
        for item in sorted(items, key=lambda x: x["Expiry date"]):
            status = batch_status(item, days)
            if status == "CLOSED" and not include_closed:
                continue
            visible.append({
                "Batch ID": item["Batch ID"], "Ingredient": item["Ingredient"],
                "Category": item["Category"], "Remaining": item["Qty remaining"],
                "Unit": item["Unit"], "Expiry date": item["Expiry date"],
                "Days left": (item["Expiry date"] - date.today()).days,
                "Status": status, "Supplier": item["Supplier"],
            })
        if visible:
            frame = pd.DataFrame(visible)
            def shade(row):
                return [f"background-color: {COLOURS[row['Status']]}"] * len(row)
            st.dataframe(frame.style.apply(shade, axis=1), hide_index=True, use_container_width=True)
        else:
            st.info("No inventory to display yet. Add your first batch in Add stock.")
        st.download_button(
            "⬇️ Download inventory Excel",
            data=DATA_FILE.read_bytes(),
            file_name="restaurant_inventory.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        st.caption("Red = past use-by/expiry; amber = approaching date or past best-before (quality review); green = later; grey = used/discarded. Always verify the physical label.")

    with intake:
        st.subheader("Record a new delivery")
        image = st.file_uploader("Upload a photo of the printed expiry label (optional)", type=["png", "jpg", "jpeg"])
        suggested_dates = []
        if image is not None:
            try:
                recognised, suggested_dates = photo_date_suggestions(image.getvalue())
                with st.expander("Text detected in photo"):
                    st.code(recognised or "No readable text found")
                if suggested_dates:
                    st.info("Possible dates from photo: " + ", ".join(d.isoformat() for d in suggested_dates))
                else:
                    st.warning("No full date was detected. Enter it manually.")
            except Exception as exc:
                st.warning(f"Photo recognition unavailable: {exc}. You can enter the date manually.")

        with st.form("new_batch", clear_on_submit=True):
            name = st.text_input("Ingredient / product name *", placeholder="e.g. Paneer")
            category = st.selectbox("Category", ["Dairy", "Vegetables", "Meat / seafood", "Bakery", "Dry goods", "Other"])
            quantity = st.number_input("Quantity received *", min_value=0.01, value=1.0, step=0.5)
            unit = st.selectbox("Unit", ["kg", "g", "litres", "ml", "pieces", "packs"])
            received = st.date_input("Date received", value=date.today())
            if suggested_dates:
                choice = st.selectbox("Choose a suggested date (or edit below)", suggested_dates)
            else:
                choice = date.today()
            expiry = st.date_input("Expiry / best-before date *", value=choice)
            date_type = st.selectbox("What does the package say?", ["Use by / Expiry", "Best before", "Unknown"])
            supplier = st.text_input("Supplier")
            storage = st.text_input("Storage location", placeholder="e.g. Walk-in refrigerator")
            notes = st.text_area("Notes", height=70)
            confirmed = st.checkbox("I checked the expiry date against the actual product label")
            submit = st.form_submit_button("Save batch to Excel")
        if submit:
            if not confirmed:
                st.error("Please check the physical package and tick the confirmation checkbox.")
            else:
                try:
                    batch_id = add_batch(name, category, received, expiry, date_type,
                                         quantity, unit, supplier, storage, notes)
                    st.success(f"Saved! Batch ID: {batch_id}. The Excel workbook has been updated.")
                except ValueError as exc:
                    st.error(str(exc))

    with usage:
        st.subheader("Record used or discarded ingredients")
        available = [item for item in items if batch_status(item, days) != "CLOSED"]
        if not available:
            st.info("No active batches available.")
        else:
            mapping = {f"{i['Ingredient']} | {i['Batch ID']} | {i['Qty remaining']:g} {i['Unit']} left": i for i in available}
            with st.form("movement"):
                selected_label = st.selectbox("Choose batch", list(mapping))
                selected = mapping[selected_label]
                action = st.selectbox("What happened?", ["Used", "Discarded"])
                taken = st.number_input("Quantity", min_value=0.001, max_value=float(selected["Qty remaining"]), value=min(1.0, float(selected["Qty remaining"])), step=0.1)
                note = st.text_input("Reason / note", placeholder="e.g. used for lunch service")
                record = st.form_submit_button("Save movement")
            if record:
                try:
                    remaining = record_movement(selected["Batch ID"], action, taken, note)
                    st.success(f"Movement saved. {remaining:g} {selected['Unit']} remain. Refresh the page to see updated totals.")
                except ValueError as exc:
                    st.error(str(exc))

    with settings:
        st.subheader("Alerts and notification settings")
        new_days = st.number_input("Warn this many days before expiry", min_value=0, max_value=90, value=days, step=1)
        if st.button("Save alert window"):
            set_alert_days(int(new_days))
            st.success("Alert window updated. Refresh the page to see updated colours.")
        st.markdown("*Optional email alerts:* Create a .env file using .env.example and enter your SMTP details. Never upload .env to GitHub.")
        if st.button("Send alert email now"):
            try:
                st.success(send_email_alerts())
            except Exception as exc:
                st.error(f"Email not sent: {exc}")
        st.code("python app.py --check-alerts", language="bash")
        st.caption("For daily email notifications, schedule this command once per day with your computer's scheduler. Alerts are not sent automatically while the computer is switched off.")
        st.warning("This is an inventory reminder, not a food-safety certification. Best-before and use-by dates are different. Never use a photographed date without checking the original label.")


if __name__ == "__main__":
    if "--check-alerts" in sys.argv:
        try:
            print(send_email_alerts())
        except Exception as error:
            print(f"Alert check failed: {error}", file=sys.stderr)
            sys.exit(1)
    else:
        main()